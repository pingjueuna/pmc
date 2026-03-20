"""Excel 내보내기 서비스"""
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import List


COLUMNS = [
    ("역량", 15),
    ("역량NO", 10),
    ("세부역량 수행목표", 30),
    ("난이도", 8),
    ("방법론", 10),
    ("문항 제목", 25),
    ("문항 유형", 10),
    ("예상 오답율", 10),
    ("오답유형Tag", 12),
    ("문제", 50),
    ("보기1", 30),
    ("보기2", 30),
    ("보기3", 30),
    ("보기4", 30),
    ("보기5", 30),
    ("정답 번호", 10),
    ("정답 해설", 40),
    ("출제 기준 교육 모듈", 25),
    ("참고 문서", 30),
    ("중복 검증 결과", 25),
]

FIELD_MAP = {
    "역량": "competency",
    "역량NO": "competency_no",
    "세부역량 수행목표": "sub_goal",
    "난이도": "difficulty",
    "방법론": "methodology",
    "문항 제목": "title",
    "문항 유형": "question_type",
    "예상 오답율": "expected_wrong_rate",
    "오답유형Tag": "wrong_type_tag",
    "문제": "question_text",
    "보기1": "choice1",
    "보기2": "choice2",
    "보기3": "choice3",
    "보기4": "choice4",
    "보기5": "choice5",
    "정답 번호": "answer",
    "정답 해설": "explanation",
    "출제 기준 교육 모듈": "education_module",
    "참고 문서": "reference",
    "중복 검증 결과": "validation_result",
}


def export_questions_to_excel(questions: List[dict]) -> bytes:
    """문항 목록을 Excel 파일로 변환하여 bytes 반환"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "생성 문항"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="top", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )

    # 헤더 작성
    for col_idx, (col_name, col_width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    # 난이도별 색상
    difficulty_colors = {
        "하": "E2EFDA",
        "중": "FFF2CC",
        "상": "FCE4D6",
    }

    # 데이터 작성
    for row_idx, question in enumerate(questions, 2):
        difficulty = question.get("difficulty", "")
        row_color = difficulty_colors.get(difficulty, "FFFFFF")
        row_fill = PatternFill(start_color=row_color, end_color=row_color, fill_type="solid")

        for col_idx, (col_name, _) in enumerate(COLUMNS, 1):
            field_key = FIELD_MAP.get(col_name, "")
            value = question.get(field_key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value) if value else "")
            cell.fill = row_fill
            cell.border = thin_border

            if col_idx in [1, 2, 4, 5, 7, 8, 9, 16]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

        ws.row_dimensions[row_idx].height = 80

    # 탭 구분 텍스트 시트 추가 (엑셀 붙여넣기용)
    ws2 = wb.create_sheet("붙여넣기용(탭구분)")
    header_row = "\t".join([col_name for col_name, _ in COLUMNS])
    ws2.cell(row=1, column=1, value=header_row)

    for row_idx, question in enumerate(questions, 2):
        values = []
        for col_name, _ in COLUMNS:
            field_key = FIELD_MAP.get(col_name, "")
            values.append(str(question.get(field_key, "") or ""))
        ws2.cell(row=row_idx, column=1, value="\t".join(values))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()
